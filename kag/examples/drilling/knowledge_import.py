# Copyright 2026 向金涛
# Affiliation: 吉林大学
"""核对本地数据，按确定字段建图与文本索引，不让大模型重新猜已有结构。"""

import argparse
import csv
import hashlib
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent
SOURCES = ROOT / "knowledge/sources.json"
OUTPUT = ROOT / "data/knowledge"


def save_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def digest(path):
    result = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            result.update(block)
    return result.hexdigest()


def local_file(root, relative):
    path = (root / relative).resolve()
    if not path.is_relative_to(root.resolve()) or not path.is_file():
        raise ValueError(f"源文件不存在或超出数据目录：{relative}")
    return path


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def workbook_profile(path, target):
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    formulas = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = book.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        header = list(next(iterator))
        original = formulas.worksheets[0].iter_rows(values_only=True)
        next(original)
        while header and header[-1] is None:
            header.pop()
        if not header or None in header or len(set(header)) != len(header) or target not in header:
            raise ValueError("工作簿列名不完整或缺少目标列")
        missing, labels, formula_columns, seen = Counter(), Counter(), Counter(), set()
        count = duplicates = zeros = positives = formula_cells = 0
        low = high = None
        for row, raw in zip(iterator, original):
            if all(value is None for value in row):
                continue
            if any(value is not None for value in row[len(header):]):
                raise ValueError("存在无列名的数据")
            values = tuple(row[:len(header)])
            for key, cached, value in zip(header, values, raw):
                if isinstance(value, str) and value.startswith("="):
                    formula_cells += 1
                    formula_columns[key] += 1
                    if cached is None:
                        raise ValueError("源工作簿公式缺少缓存值，不能统计")
            count += 1
            duplicates += values in seen
            seen.add(values)
            missing.update(key for key, value in zip(header, values) if value is None)
            value = values[header.index(target)]
            if value is None:
                raise ValueError("目标值缺失")
            if isinstance(value, str):
                labels[value] += 1
            else:
                low = value if low is None else min(low, value)
                high = value if high is None else max(high, value)
                zeros += value == 0
                positives += value > 0
        return {"sheet": sheet.title, "rows": count, "columns": header, "target": target,
                "missing_by_column": {key: missing[key] for key in header}, "duplicate_rows": duplicates,
                "label_counts": dict(labels), "target_range": [low, high] if low is not None else None,
                "target_zero": zeros, "target_positive": positives,
                "formula_cells": formula_cells, "formula_by_column": dict(formula_columns),
                "formula_policy": "读取发布者保存的公式缓存，未重新计算或修改原工作簿"}
    finally:
        book.close()
        formulas.close()


def prepare(root):
    import pyarrow.parquet as pq

    config = json.loads(SOURCES.read_text(encoding="utf-8"))
    split_path = local_file(root, config["split_manifest"])
    splits = {(row["wellbore_id"], row["split"]) for row in read_csv(split_path)}
    if {role for well, role in splits if well == config["development_well"]} != {"development"}:
        raise ValueError("开发井的划分不明确，停止导入")
    manifest = {row["local_path"].replace("\\", "/"): row for row in
                read_csv(local_file(root, "01_inventory/download_manifest.csv")) if row["local_path"]}
    evidence_path = local_file(root, "01_inventory/well_event_screening_evidence.csv")
    screening = {}
    for row in read_csv(evidence_path):
        if row["wellbore_id"] == config["development_well"] and row["source_kind"] == "activity":
            key = (row["doc_name"], row["activity_start"], row["activity_end"])
            screening.setdefault(key, set()).add(row["category"])
    datasets, records = [], []
    for source in config["datasets"]:
        path = local_file(root, source["path"])
        expected = source["sha256"]
        if source["path"] in manifest and manifest[source["path"]]["sha256"] != expected:
            raise ValueError("下载清单与项目固定指纹不一致")
        actual = digest(path)
        if not expected or actual != expected:
            raise ValueError(f"源文件指纹不匹配：{source['id']}")
        dataset = {**source, "sha256": actual}
        if source.get("target"):
            dataset["profile"] = workbook_profile(path, source["target"])
        else:
            all_reports = pq.read_table(path).to_pylist()
            reports = [report for report in all_reports if report["nameWellbore"] == config["development_well"]]
            if not reports:
                raise ValueError("开发井没有日报")
            if len({report["docName"] for report in reports}) != len(reports):
                raise ValueError("日报标识重复，停止覆盖")
            for report in reports:
                report_id = "volve:" + report["docName"]
                summaries = [item["sum24Hr"] for item in report.get("statusInfo") or []
                             if item.get("sum24Hr") not in {None, "", "None"}]
                common = {"well": report["nameWellbore"], "dataset": source["id"], "source_url": source["source_url"],
                          "doc_name": report["docName"], "report_start": report["dTimStart"], "report_end": report["dTimEnd"]}
                records.append({**common, "id": report_id, "kind": "report", "anchor": report["docName"] + ":statusInfo.sum24Hr",
                                "text": "\n".join(summaries) or "日报未提供非空的24小时摘要，详见其作业条目。",
                                "start": report["dTimStart"], "end": report["dTimEnd"], "categories": [], "state": "", "detail": ""})
                for index, activity in enumerate(report.get("activity") or []):
                    key = (report["docName"], activity["dTimStart"], activity["dTimEnd"])
                    records.append({**common, "id": f"{report_id}:activity:{index}", "kind": "activity", "parent": report_id,
                                    "anchor": f"{report['docName']}:activity[{index}]", "text": activity.get("comments") or "未记录作业正文",
                                    "start": activity["dTimStart"], "end": activity["dTimEnd"],
                                    "categories": sorted(screening.get(key, [])), "state": activity.get("state") or "",
                                    "detail": activity.get("stateDetailActivity") or ""})
            dataset["profile"] = {"source_reports": len(all_reports), "imported_reports": len(reports),
                                  "imported_activities": sum(record["kind"] == "activity" for record in records),
                                  "well": config["development_well"], "split": "development"}
        datasets.append(dataset)
    inventory = [{"file": path.name, "rows": len(read_csv(path)), "sha256": digest(path)}
                 for path in sorted((root / "01_inventory").glob("*.csv"))]
    return {"datasets": datasets, "records": records, "inventory": inventory,
            "lineage": {"screening_sha256": digest(evidence_path), "split_sha256": digest(split_path)},
            "excluded_splits": sorted({role for _, role in splits if role != "development"}),
            "scope": "开发井的回溯式日报知识；两份井漏表的统计及使用边界。原始数值表和时序文件留在源目录，未导入验证/测试案例及实验结果。"}


def text_parts(text, size=1000):
    while text:
        end = min(size, len(text))
        if end < len(text):
            boundary = text.rfind("\n", size // 2, end)
            if boundary >= 0:
                end = boundary + 1
        yield text[:end]
        text = text[end:]


def graph_data(bundle):
    nodes, edges, catalog = [], [], []
    def node(key, name, label, **properties):
        item = {"id": key, "name": name, "label": label, "properties": properties}
        nodes.append(item)
        return item
    def edge(start, relation, end):
        edges.append({"id": "", "from": start["id"], "fromType": start["label"], "to": end["id"],
                      "toType": end["label"], "label": relation, "properties": {}})
    def chunks(owner, title, content, url, kind):
        # 每块重复来源和定位，正文按字符完整切分，避免静默截断长作业记录。
        for index, part in enumerate(text_parts(content)):
            key = owner["id"] + f":chunk:{index}"
            body = f"{title}\n来源：{url}\n{part}"
            chunk = node(key, title, "Chunk", content=body)
            edge(owner, "source", chunk)
            catalog.append({"id": key, "title": title, "file": key, "kind": kind, "url": url, "content": body, "page": ""})
    datasets = {}
    for dataset in bundle["datasets"]:
        current = node(dataset["id"], dataset["name"], "KnowledgeDataset", sourceUrl=dataset["source_url"],
                       fileHash=dataset["sha256"], attribution=dataset["attribution"], license=dataset["license"],
                       profile=json.dumps(dataset["profile"], ensure_ascii=False), scope=dataset["limitations"])
        datasets[dataset["id"]] = current
        content = f"任务：{dataset['task']}\n署名：{dataset['attribution']}\n许可：{dataset['license']}\n限制：{dataset['limitations']}\n实测统计：{json.dumps(dataset['profile'], ensure_ascii=False, indent=2)}"
        chunks(current, dataset["name"], content, dataset["source_url"], "数据集说明与实测统计")
    wells, record_nodes = {}, {}
    for record in bundle["records"]:
        if record["well"] not in wells:
            wells[record["well"]] = node(record["well"], record["well"], "Wellbore", usage="development")
        current = node(record["id"], record["anchor"], "DrillingRecord", kind=record["kind"],
                       anchor=record["anchor"], startTime=record["start"], endTime=record["end"],
                       reportStart=record["report_start"], reportEnd=record["report_end"],
                       text=record["text"], state=record["state"], stateDetail=record["detail"],
                       categories=json.dumps(record["categories"]), reviewStatus="SOURCE_RECORD_NOT_ROOT_CAUSE_LABEL")
        record_nodes[record["id"]] = current
        edge(current, "fromDataset", datasets[record["dataset"]])
        edge(current, "inWell", wells[record["well"]])
        if record.get("parent"):
            edge(current, "inReport", record_nodes[record["parent"]])
        title = f"Volve 钻井 {record['well']} {record['anchor']}"
        content = (f"发生时段：{record['start']} 至 {record['end']}；报告时段：{record['report_start']} 至 {record['report_end']}\n"
                   f"原始状态：{record['state']}；状态细项：{record['detail']}\n候选筛查类别：{','.join(record['categories']) or '未命中'}。候选命中不构成确认异常或根因。\n{record['text']}")
        chunks(current, title, content, record["source_url"], "开发井日报原文" if record["kind"] == "report" else "开发井作业原文")
    return {"resultNodes": nodes, "resultEdges": edges}, catalog


def verify_graph(graph, connection):
    """从 OpenSPG 回读节点及所有关系，写入接口返回成功不作为唯一验收依据。"""
    from kag.common.tools.graph_api.impl.openspg_graph_api import OpenSPGGraphApi
    api = OpenSPGGraphApi(**connection)
    def decode(value):
        return json.loads(value) if isinstance(value, str) and value.startswith('"') else value
    for label in {node["label"] for node in graph["resultNodes"]}:
        expected = {node["id"] for node in graph["resultNodes"] if node["label"] == label}
        actual = api.execute_dsl(f"MATCH (n:`Drilling.{label}`) RETURN n.id")
        if not expected <= {decode(row[0]) for row in actual.data}:
            raise ValueError(f"图谱回读缺少节点：{label}")
    for relation in {edge["label"] for edge in graph["resultEdges"]}:
        expected = {(edge["from"], edge["to"]) for edge in graph["resultEdges"] if edge["label"] == relation}
        actual = api.execute_dsl(f"MATCH (n)-[:{relation}]->(m) RETURN n.id,m.id")
        if not expected <= {(decode(row[0]), decode(row[1])) for row in actual.data}:
            raise ValueError(f"图谱回读缺少关系：{relation}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path("D:/drilling_agent_data"))
    parser.add_argument("--write", action="store_true", help="同步业务Schema、向量化并写入本地OpenSPG")
    parser.add_argument("--config", type=Path, default=ROOT / "kag_config.local.yaml")
    args = parser.parse_args()
    bundle = prepare(args.root)
    graph, catalog = graph_data(bundle)
    report = {"created_at": datetime.now(timezone.utc).isoformat(), "scope": bundle["scope"],
              "datasets": bundle["datasets"], "lineage": bundle["lineage"], "inventory": bundle["inventory"],
              "excluded_splits": bundle["excluded_splits"],
              "nodes": dict(Counter(node["label"] for node in graph["resultNodes"])), "edges": len(graph["resultEdges"]),
              "status": "prepared"}
    if args.write:
        import run as cli
        from kag.builder.model.sub_graph import SubGraph
        from kag.interface import VectorizerABC, SinkWriterABC
        config = cli.load_config(args.config)
        cli.validate_config(config)
        cli.init_runtime(config)
        cli.commit_schema(config)
        vectorizer = VectorizerABC.from_config({"type": "batch_vectorizer", "vectorize_model": config["vectorize_model"], "batch_size": 16})
        writer = SinkWriterABC.from_config({"type": "kg_writer"})
        # 稳定ID支持幂等重跑；最后才发布成功清单，失败时不展示“已入库”。
        for start in range(0, len(graph["resultNodes"]), 64):
            part = SubGraph.from_dict({"resultNodes": graph["resultNodes"][start:start + 64], "resultEdges": []})
            for result in vectorizer.invoke(part, write_ckpt=False):
                writer.invoke(result, write_ckpt=False)
            print(f"nodes {min(start + 64, len(graph['resultNodes']))}/{len(graph['resultNodes'])}", flush=True)
        writer.invoke(SubGraph.from_dict({"resultNodes": [], "resultEdges": graph["resultEdges"]}), write_ckpt=False)
        verify_graph(graph, {"host_addr": config["project"]["host_addr"], "project_id": config["project"]["id"]})
        report["status"] = "written"
        save_json(OUTPUT / "catalog.json", catalog)
        save_json(OUTPUT / "status.json", report)
    output = ROOT / "results/knowledge" / ("import.json" if args.write else "prepared.json")
    save_json(output, report)
    print(json.dumps({"status": report["status"], "nodes": report["nodes"], "edges": report["edges"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
